"""Base consolidada em Excel: uma ação por linha, pronta para filtro ou Power BI."""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from modules import qualidade, regras

ROXO = "5E22F3"


def _formatar(ws, larguras):
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=ROXO)
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def gerar_base(df: pd.DataFrame, hoje: date) -> bytes:
    acoes = pd.DataFrame({
        "Frente": df["frente"], "Entrega": df["entrega"], "Subgrupo": df["subgrupo"], "Ação": df["acao"],
        "Responsável": df["responsavel"], "Status": df["status"], "Peso": df["peso"], "% (pelo status)": df["pct"],
        "Data prevista/conclusão": df["data"], "Sinal": df["rotulo_sinal"].str.split(" ", n=1).str[1],
        "Dias de atraso": df["dias_atraso"], "Aguardando terceiro": df["aguardando"].map({True: "Sim", False: ""}),
        "Área de apoio": df["externo"].map({True: "Sim", False: ""}), "Último comentário": df["comentario"],
        "Aba de origem": df["aba"], "Linha de origem": df["linha"],
    })
    if "vezes_prazo_alterado" in df.columns:  # com histórico gravado
        acoes.insert(9, "Prazo original", df["prazo_original"].values)
        acoes.insert(10, "Prazo adiado (vezes)", df["vezes_prazo_alterado"].values)
        acoes.insert(11, "Entrou em", df["incluida_em"].values)
        acoes.insert(12, "Está na planilha atual", df["ativa"].map({True: "Sim", False: "Não"}).values)
    ent = regras.resumo_entregas(df)
    entregas = pd.DataFrame({
        "Frente": ent["frente"], "Entrega": ent["entrega"], "Farol": ent["rotulo_farol"].str.split(" ", n=1).str[1],
        "Ações": ent["acoes"], "Em aberto": ent["abertas"], "Atrasadas": ent["atrasadas"],
        "Avanço": ent["avanco"], "Prazo final": ent["prazo_final"],
    }) if len(ent) else pd.DataFrame()
    pend = []
    for c in qualidade.checar(df):
        for r in c["acoes"].itertuples():
            pend.append({"Checagem": c["checagem"], "Frente": r.frente, "Entrega": r.entrega, "Ação": r.acao,
                         "Aba de origem": r.aba, "Linha de origem": r.linha})
    pend = pd.DataFrame(pend, columns=["Checagem", "Frente", "Entrega", "Ação", "Aba de origem", "Linha de origem"])

    saida = io.BytesIO()
    with pd.ExcelWriter(saida, engine="openpyxl") as xw:
        acoes.to_excel(xw, sheet_name="Ações", index=False)
        larguras = {"Frente": 22, "Entrega": 30, "Subgrupo": 18, "Ação": 45, "Responsável": 18,
                    "Último comentário": 50, "Aba de origem": 18}
        _formatar(xw.sheets["Ações"], [larguras.get(c, 13) for c in acoes.columns])
        for linha in xw.sheets["Ações"].iter_rows(min_row=2):
            for cel in linha:
                cab = xw.sheets["Ações"].cell(1, cel.column).value
                if cab == "% (pelo status)":
                    cel.number_format = "0%"
                elif cab in ("Data prevista/conclusão", "Prazo original", "Entrou em"):
                    cel.number_format = "DD/MM/YYYY"
        if len(entregas):
            entregas.to_excel(xw, sheet_name="Entregas", index=False)
            _formatar(xw.sheets["Entregas"], [22, 40, 12, 8, 10, 10, 9, 12])
            for linha in xw.sheets["Entregas"].iter_rows(min_row=2, min_col=7, max_col=8):
                linha[0].number_format = "0%"
                linha[1].number_format = "DD/MM/YYYY"
        pend.to_excel(xw, sheet_name="Qualidade dos dados", index=False)
        _formatar(xw.sheets["Qualidade dos dados"], [38, 22, 30, 45, 18, 9])
        info = xw.book.create_sheet("Sobre")
        info["A1"] = "Gerado pelo app Gestão de Demandas"
        info["A1"].font = Font(bold=True, color=ROXO, size=13)
        info["A2"] = f"Posição em {hoje:%d/%m/%Y}"
        info["A3"] = "Sinal: Atrasada = prazo < data de posição e não concluída; " \
                     f"Vence = prazo nos próximos {regras.JANELA_DIAS} dias."
        info["A4"] = "Avanço = Σ(peso × %) ÷ Σ peso, com % pelo status (0 / 50 / 100%)."
    return saida.getvalue()
