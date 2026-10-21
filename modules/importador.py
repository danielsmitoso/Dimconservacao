"""Leitura da planilha de demandas da equipe.

As colunas são encontradas PELO NOME do cabeçalho (não pela posição), para
que inserir ou mover uma coluna na planilha não quebre a leitura.

Modelo atual (rev.01, out/2026) -- uma aba por FRENTE:
    B Avanço | C Título/Tarefa | D Tarefa | E Responsável | F Status | G Peso |
    H % | I Prazo/Conclusão | J em diante: comentários por semana
- coluna "Título/Tarefa" diz o tipo da linha: "Título" = ENTREGA,
  "Tarefa" = AÇÃO ("Subtítulo"/"Subgrupo", se usado, = SUBGRUPO);
- "%" é lido da planilha: vale o % digitado (ex.: 80% em andamento);
  sem valor, sai do status (0 / 50 / 100%); Concluído é sempre 100%.

Modelo antigo (até 19/10/2026), ainda aceito -- sem a coluna de tipo:
    C Tarefa | D Responsável | E Status | F Peso | G % | H Data | I+ semanas
  e o tipo da linha vem da formatação: negrito com fundo = entrega,
  negrito sem fundo = subgrupo, demais = ação.

As fórmulas de avanço da planilha NÃO são usadas: o app recalcula a partir
das ações.
"""
from __future__ import annotations

import io
import re
import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
import pandas as pd

LINHAS_BUSCA_CABECALHO = 20

# Linhas de modelo que vêm vazias na planilha ("AÇÕES" sem responsável etc.)
NOMES_MODELO = {"AÇÕES", "TAREFA PRINCIPAL"}
STATUS_VALIDOS = ("À iniciar", "Em andamento", "Concluído")
SEM_STATUS = "Sem status"
SEM_ENTREGA = "(sem entrega)"

COLUNAS = ["frente", "entrega", "subgrupo", "acao", "responsavel", "status", "peso", "pct",
           "data", "comentario", "historico", "nota", "aba", "linha"]

MODELO_ATUAL, MODELO_ANTIGO = "atual (com coluna Título/Tarefa)", "antigo (tipo pela formatação)"


@dataclass
class Resultado:
    acoes: pd.DataFrame
    abas_lidas: list[str] = field(default_factory=list)
    abas_vazias: list[str] = field(default_factory=list)   # só têm o modelo, sem ação real
    abas_ignoradas: list[str] = field(default_factory=list)  # sem cabeçalho "Tarefa" (ex.: painel de indicadores)
    avisos: list[str] = field(default_factory=list)          # problemas de estrutura
    modelos: dict[str, str] = field(default_factory=dict)    # aba -> modelo reconhecido


def _texto(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _norm(v) -> str:
    """Minúsculo, sem acento, espaços simples (para comparar cabeçalhos e tipos)."""
    t = unicodedata.normalize("NFKD", _texto(v)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), formato).date()
            except ValueError:
                pass
    return None


def _numero(v) -> float | None:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    return float(v)


def _tem_fundo(celula) -> bool:
    preench = celula.fill
    if preench is None or not preench.fill_type:
        return False
    cor = preench.fgColor
    if cor is None:
        return False
    if cor.type == "rgb":
        return cor.rgb not in (None, "00000000", "FFFFFFFF")
    return True  # cor de tema/indexada: considera preenchido


def _normalizar_status(v) -> str:
    t = _texto(v)
    if not t:
        return SEM_STATUS
    for s in STATUS_VALIDOS:
        if _norm(t) == _norm(s):
            return s
    return t


def _pct(status: str) -> float:
    """% padrão pelo status (quando a planilha não traz o %)."""
    return {"Concluído": 1.0, "Em andamento": 0.5}.get(status, 0.0)


def pct_da_linha(status: str, pct_planilha) -> float:
    """Concluído = 100%; senão vale o % digitado (0 a 100%); sem %, o padrão do status."""
    if status == "Concluído":
        return 1.0
    v = _numero(pct_planilha)
    if v is None:
        return _pct(status)
    if v > 1:  # digitaram 80 em vez de 80%
        v = v / 100
    return min(max(v, 0.0), 1.0)


# ---------------------------------------------------------------------------
# Cabeçalho
# ---------------------------------------------------------------------------
def _papel(cabecalho: str) -> str | None:
    """Que informação a coluna guarda, pelo texto do cabeçalho."""
    h = _norm(cabecalho)
    if not h:
        return None
    if h in ("titulo/tarefa", "tipo", "tipo de linha", "titulo / tarefa"):
        return "tipo"
    if h in ("tarefa", "acao", "acoes", "atividade"):
        return "nome"
    if h.startswith("responsavel"):
        return "resp"
    if h == "status":
        return "status"
    if h == "peso":
        return "peso"
    if h in ("%", "% concluido", "percentual"):
        return "pct"
    if h.startswith("avanco"):
        return "avanco"
    if h.startswith("prazo") or h.startswith("data"):
        return "data"
    return None


def _cabecalho(ws) -> tuple[int, dict, list] | None:
    """Linha do cabeçalho, {papel: coluna} e as colunas de semana [(col, rótulo)]."""
    for r in range(1, min(LINHAS_BUSCA_CABECALHO, ws.max_row) + 1):
        papeis: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            p = _papel(ws.cell(r, c).value)
            if p and p not in papeis:
                papeis[p] = c
        if "nome" in papeis and "status" in papeis:
            semanas = []
            inicio = papeis.get("data", max(papeis.values())) + 1
            for c in range(inicio, ws.max_column + 1):
                v = ws.cell(r, c).value
                d = _data(v)
                if d is not None:
                    semanas.append((c, f"{d:%d/%m/%Y}"))
                elif _texto(v):
                    semanas.append((c, _texto(v)))
            return r, papeis, semanas
    return None


def _tipo_da_linha(ws, r: int, papeis: dict) -> str:
    """'entrega', 'subgrupo' ou 'acao'."""
    if "tipo" in papeis:
        t = _norm(ws.cell(r, papeis["tipo"]).value)
        if t.startswith("titulo") or t.startswith("entrega"):
            return "entrega"
        if t.startswith("subtitulo") or t.startswith("subgrupo"):
            return "subgrupo"
        if t.startswith("tarefa") or t.startswith("acao"):
            return "acao"
    celula = ws.cell(r, papeis["nome"])  # sem coluna de tipo (modelo antigo) ou tipo em branco
    negrito = bool(celula.font and celula.font.bold)
    if negrito and _tem_fundo(celula):
        return "entrega"
    return "subgrupo" if negrito else "acao"


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------
def _ler_aba(ws, res: Resultado) -> list[dict]:
    achado = _cabecalho(ws)
    if achado is None:
        res.abas_ignoradas.append(ws.title)
        return []
    cab, papeis, semanas = achado
    col = lambda r, papel: ws.cell(r, papeis[papel]).value if papel in papeis else None  # noqa: E731

    linhas, avisos = [], []
    entrega, subgrupo = None, None
    sem_nome = 0
    for r in range(cab + 1, ws.max_row + 1):
        nome = _texto(col(r, "nome"))
        tipo = _tipo_da_linha(ws, r, papeis)
        if not nome:
            if tipo == "acao" and (_texto(col(r, "resp")) or _texto(col(r, "status"))):
                sem_nome += 1  # tem responsável/status mas não tem o nome da tarefa
            continue
        if tipo == "entrega":
            entrega, subgrupo = nome, None
            if nome.upper() == "TAREFA PRINCIPAL":
                avisos.append(f"{ws.title}: entrega sem nome (\"TAREFA PRINCIPAL\") na linha {r}")
            d = _data(col(r, "data"))
            if d and d.year < 2025:
                avisos.append(f"{ws.title} › {nome}: data da entrega {d:%d/%m/%Y} parece digitada errada (linha {r})")
            continue
        if tipo == "subgrupo":
            subgrupo = nome
            continue

        resp = _texto(col(r, "resp"))
        if nome.upper() in NOMES_MODELO and not resp:
            continue
        status = _normalizar_status(col(r, "status"))
        historico = [(sem, _texto(ws.cell(r, c).value)) for c, sem in semanas
                     if _texto(ws.cell(r, c).value) and _norm(ws.cell(r, c).value) != "status/comentario"]
        nota = col(r, "avanco")
        d = _data(col(r, "data"))
        if d and d.year < 2025:
            avisos.append(f"{ws.title} › {nome}: prazo {d:%d/%m/%Y} parece digitado errado (linha {r})")
        linhas.append({
            "frente": ws.title.strip(),
            "entrega": entrega or SEM_ENTREGA,
            "subgrupo": subgrupo or "",
            "acao": nome,
            "responsavel": resp,
            "status": status,
            "peso": _numero(col(r, "peso")),
            "pct": pct_da_linha(status, col(r, "pct")),
            "data": d,
            "comentario": historico[-1][1] if historico else "",
            "historico": historico,
            "nota": nota.strip() if isinstance(nota, str) and not nota.startswith("#") else "",
            "aba": ws.title,
            "linha": r,
        })
    if sem_nome:
        avisos.append(f"{ws.title}: {sem_nome} linha(s) de tarefa com responsável/status mas sem o nome da tarefa (ignoradas)")
    if linhas:
        res.abas_lidas.append(ws.title)
        res.modelos[ws.title] = MODELO_ATUAL if "tipo" in papeis else MODELO_ANTIGO
        res.avisos.extend(avisos)  # aba só com o modelo não gera aviso
    else:
        res.abas_vazias.append(ws.title)
    return linhas


def ler_planilha(arquivo) -> Resultado:
    """Lê a planilha (caminho, bytes ou arquivo aberto) e devolve uma ação por linha."""
    if isinstance(arquivo, (bytes, bytearray)):
        arquivo = io.BytesIO(arquivo)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # extensões de formatação condicional não suportadas
        wb = openpyxl.load_workbook(arquivo, data_only=True)
    res = Resultado(acoes=pd.DataFrame(columns=COLUNAS))
    linhas = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        linhas.extend(_ler_aba(ws, res))
    if linhas:
        df = pd.DataFrame(linhas, columns=COLUNAS)
        df.insert(0, "id", range(1, len(df) + 1))
        res.acoes = df
    return res
