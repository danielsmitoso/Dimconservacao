"""Relatório executivo semanal em PowerPoint (mesmo padrão visual do DimConservação)."""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from modules import regras

ROXO_PRIMARIO = RGBColor(0x5E, 0x22, 0xF3)
ROXO_ESCURO = RGBColor(0x3D, 0x1A, 0x9E)
LAVANDA = RGBColor(0xED, 0xE8, 0xFB)
CARD = RGBColor(0xF7, 0xF4, 0xFE)
BORDA = RGBColor(0xE4, 0xDA, 0xFB)
TINTA = RGBColor(0x1A, 0x1A, 0x2E)
CINZA = RGBColor(0x55, 0x55, 0x55)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MAX_LINHAS_TABELA = 12

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"
LOGO_BRANCA = ASSETS_DIR / "logo_motiva_branca.png"
LOGO_ROXA = ASSETS_DIR / "logo_motiva_roxa.png"


def _slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _titulo(slide, texto, subtitulo=None):
    if LOGO_ROXA.exists():
        slide.shapes.add_picture(str(LOGO_ROXA), SLIDE_W - Inches(1.7), Inches(0.35), height=Inches(0.35))
    tf = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(10.9), Inches(0.9)).text_frame
    tf.text = texto
    p = tf.paragraphs[0]
    p.font.size, p.font.bold, p.font.color.rgb = Pt(26), True, ROXO_ESCURO
    if subtitulo:
        p2 = tf.add_paragraph()
        p2.text = subtitulo
        p2.font.size, p2.font.color.rgb = Pt(13), CINZA


def _capa(prs, titulo, subtitulo):
    slide = _slide(prs)
    fundo = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    fundo.fill.solid()
    fundo.fill.fore_color.rgb = ROXO_PRIMARIO
    fundo.line.fill.background()
    fundo.shadow.inherit = False
    if LOGO_BRANCA.exists():
        slide.shapes.add_picture(str(LOGO_BRANCA), Inches(0.8), Inches(0.6), height=Inches(0.5))
    for texto, top, tam, cor, negrito in ((titulo, 2.8, 40, BRANCO, True), (subtitulo, 3.9, 18, LAVANDA, False)):
        tf = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(1)).text_frame
        tf.text = texto
        tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(tam), negrito
        tf.paragraphs[0].font.color.rgb = cor


def _card(slide, left, top, rotulo, valor, detalhe):
    caixa = slide.shapes.add_shape(5, left, top, Inches(2.9), Inches(1.35))
    caixa.fill.solid()
    caixa.fill.fore_color.rgb = CARD
    caixa.line.color.rgb = BORDA
    caixa.shadow.inherit = False
    tf = caixa.text_frame
    tf.word_wrap = True
    for i, (txt, tam, cor, negrito) in enumerate(((rotulo, 12, ROXO_PRIMARIO, True), (str(valor), 28, TINTA, True),
                                                  (detalhe, 10, CINZA, False))):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size, p.font.color.rgb, p.font.bold = Pt(tam), cor, negrito


def _tabela(slide, top, dados, larguras):
    linhas, colunas = len(dados), len(dados[0])
    tabela = slide.shapes.add_table(linhas, colunas, Inches(0.6), top, Inches(12.1), Inches(0.35 * linhas)).table
    for i, w in enumerate(larguras):
        tabela.columns[i].width = Inches(w)
    for i, linha in enumerate(dados):
        for j, valor in enumerate(linha):
            cel = tabela.cell(i, j)
            cel.text = str(valor)
            p = cel.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            cel.fill.solid()
            if i == 0:
                p.font.bold, p.font.color.rgb = True, BRANCO
                cel.fill.fore_color.rgb = ROXO_PRIMARIO
            else:
                cel.fill.fore_color.rgb = BRANCO if i % 2 else LAVANDA


def _bullets(slide, top, itens, tamanho=14):
    tf = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.1), Inches(5.5)).text_frame
    tf.word_wrap = True
    for i, txt in enumerate(itens):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + txt.replace("**", "")
        p.font.size, p.font.color.rgb = Pt(tamanho), TINTA
        p.space_after = Pt(6)


def _fmt(d) -> str:
    return d.strftime("%d/%m") if d is not None and not pd.isna(d) else "—"


def _pct(v) -> str:
    return f"{v:.0%}"


def _tabela_paginada(prs, titulo, subtitulo, cabecalho, linhas, larguras):
    if not linhas:
        return
    paginas = [linhas[i:i + MAX_LINHAS_TABELA] for i in range(0, len(linhas), MAX_LINHAS_TABELA)]
    for n, pagina in enumerate(paginas, 1):
        slide = _slide(prs)
        sufixo = f" ({n}/{len(paginas)})" if len(paginas) > 1 else ""
        _titulo(slide, titulo + sufixo, subtitulo)
        _tabela(slide, Inches(1.5), [cabecalho] + pagina, larguras)


def gerar_relatorio(df: pd.DataFrame, nome_area: str, hoje: date) -> bytes:
    """df: ações já classificadas (regras.classificar) e filtradas como na tela."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _capa(prs, "Gestão de Demandas", f"{nome_area}  ·  posição em {hoje:%d/%m/%Y}")

    # Painel geral
    ind = regras.indicadores(df)
    slide = _slide(prs)
    _titulo(slide, "Painel geral", f"{ind['acoes']} ações · {ind['abertas']} em aberto")
    cards = [("Avanço geral", _pct(ind["avanco"]), "ponderado pelo peso"),
             ("Concluídas", ind["concluidas"], f"{ind['concluidas'] / max(ind['acoes'], 1):.0%} do total"),
             ("Atrasadas", ind["atrasadas"], "prazo vencido"),
             ("Vencem em 7 dias", ind["vence"], f"até {hoje + pd.Timedelta(days=regras.JANELA_DIAS):%d/%m}")]
    for i, (r, v, d) in enumerate(cards):
        _card(slide, Inches(0.6 + i * 3.07), Inches(1.4), r, v, d)
    fr = regras.resumo_frentes(df)
    dados = [["Frente", "Entregas", "Ações", "Concluídas", "Atrasadas", "Entregas críticas", "Avanço"]]
    dados += [[r.frente, r.entregas, r.acoes, r.concluidas, r.atrasadas, r.entregas_criticas, _pct(r.avanco)]
              for r in fr.itertuples()]
    _tabela(slide, Inches(3.1), dados[:MAX_LINHAS_TABELA], [3.6, 1.2, 1.2, 1.4, 1.4, 1.9, 1.4])

    # Pontos de atenção
    pontos = regras.pontos_atencao(df)
    if pontos:
        slide = _slide(prs)
        _titulo(slide, "Pontos de atenção", "gerados automaticamente a partir dos prazos e status")
        _bullets(slide, Inches(1.4), [t for _, t in pontos[:10]], 13)

    # Uma visão por frente
    ent = regras.resumo_entregas(df)
    for frente, g in ent.groupby("frente", sort=False):
        linhas = [[r.rotulo_farol, r.entrega, r.acoes, r.atrasadas, _pct(r.avanco), _fmt(r.prazo_final)]
                  for r in g.itertuples()]
        _tabela_paginada(prs, frente, f"Avanço da frente: {_pct(regras.avanco(df[df['frente'] == frente]))}",
                         ["Farol", "Entrega", "Ações", "Atrasadas", "Avanço", "Prazo final"],
                         linhas, [1.6, 5.5, 1.1, 1.3, 1.2, 1.4])

    # Pauta semanal
    blocos = regras.pauta(df, hoje)
    _tabela_paginada(prs, "Atrasado — o que destrava?", "ações com prazo vencido",
                     ["Prazo", "Ação", "Frente › Entrega", "Responsável", "Comentário"],
                     [[_fmt(r.data), r.acao, f"{r.frente} › {r.entrega}", r.responsavel, r.comentario[:80]]
                      for r in blocos["atrasado"].itertuples()], [0.9, 3.6, 3.2, 1.6, 2.8])
    _tabela_paginada(prs, "Compromissos da semana", f"prazo nos próximos {regras.JANELA_DIAS} dias",
                     ["Prazo", "Ação", "Frente › Entrega", "Responsável", "Status"],
                     [[_fmt(r.data), r.acao, f"{r.frente} › {r.entrega}", r.responsavel, r.status]
                      for r in blocos["semana"].itertuples()], [0.9, 4.2, 3.6, 1.8, 1.6])
    titulo_concl = ("Concluído desde a importação anterior" if "concluida_na_ultima" in df.columns
                    else "Concluído na última semana")
    _tabela_paginada(prs, titulo_concl, None,
                     ["Data", "Ação", "Frente › Entrega", "Responsável"],
                     [[_fmt(r.data), r.acao, f"{r.frente} › {r.entrega}", r.responsavel]
                      for r in blocos["concluido"].itertuples()], [0.9, 5.0, 4.2, 2.0])
    _tabela_paginada(prs, "Impedimentos e dependências", "aguardando terceiros ou a cargo de áreas de apoio",
                     ["Ação", "Responsável", "Comentário"],
                     [[r.acao, r.responsavel, r.comentario[:100]] for r in blocos["impedimentos"].itertuples()],
                     [4.5, 2.2, 5.4])

    saida = io.BytesIO()
    prs.save(saida)
    return saida.getvalue()
