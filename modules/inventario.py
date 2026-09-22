"""Leitura e validacao da planilha de inventario de trechos."""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from modules import parametros_padrao as pp

COLUNAS_MINIMAS = ["trecho", "km_inicial", "km_final", "area_verde_m2"]

COLUNAS_OPCIONAIS = [
    "rodovia",             # codigo da rodovia (ex: SP-330) -- permite propor equipes por rodovia
    "eps_barreira_m",      # metros lineares de EPS/barreira
    "defensa_m",           # metros lineares de defensa metalica
    "placas_qtd",          # quantidade de placas/catadioptricos
    "bueiros_qtd",         # quantidade de bueiros/dispositivos de drenagem
    "calcada_m2",          # area de calcada
    "pct_manual",          # % da area verde atendida por rocada manual (0-100)
    "pct_mecanizada",      # % da area verde atendida por rocada mecanizada (0-100)
    "equipamento_predominante",     # equipamento de rocada mecanizada mais indicado para o trecho
    "aceiro_km",                    # extensao real de aceiro (nem toda a rodovia -- exclui urbano/APP)
    "drenagem_plataforma_km",       # extensao real de drenagem de plataforma
    "drenagem_fora_plataforma_km",  # extensao real de drenagem fora de plataforma
]

COLUNAS_TEXTO = ["trecho", "rodovia", "equipamento_predominante"]

DESCRICAO_COLUNAS = {
    "trecho": "Nome/identificação do trecho (texto livre, ex: SP-330 km 12 a 17)",
    "km_inicial": "Km inicial do trecho",
    "km_final": "Km final do trecho",
    "area_verde_m2": "Área verde total do trecho, em m²",
    "rodovia": "Código da rodovia (ex: SP-330). Ajuda a propor equipes sem misturar rodovias",
    "eps_barreira_m": "Metros lineares de EPS/barreira",
    "defensa_m": "Metros lineares de defensa metálica",
    "placas_qtd": "Quantidade de placas/catadióptricos",
    "bueiros_qtd": "Quantidade de bueiros/dispositivos de drenagem",
    "calcada_m2": "Área de calçada, em m²",
    "pct_manual": "% da área verde atendida por roçada manual (0 a 100 — soma 100 com pct_mecanizada)",
    "pct_mecanizada": "% da área verde atendida por roçada mecanizada (0 a 100 — soma 100 com pct_manual)",
    "equipamento_predominante": "Equipamento de roçada mecanizada mais indicado (Trator 4.7m, Trator 1.7m, Trator com braço articulado, Eixo-zero, Spider, Robô)",
    "aceiro_km": "Extensão real de aceiro (nem sempre é a rodovia inteira — exclui trechos urbanos/APP)",
    "drenagem_plataforma_km": "Extensão real de drenagem de plataforma",
    "drenagem_fora_plataforma_km": "Extensão real de drenagem fora de plataforma",
}

_EXEMPLO_MODELO = [
    {"trecho": "SP-330 km 12 a 17", "rodovia": "SP-330", "km_inicial": 12, "km_final": 17, "area_verde_m2": 84000,
     "pct_manual": 60, "pct_mecanizada": 40, "equipamento_predominante": "Trator 1.7m",
     "eps_barreira_m": 800, "defensa_m": 1500, "placas_qtd": 22, "calcada_m2": 0,
     "aceiro_km": 4.5, "drenagem_plataforma_km": 5, "drenagem_fora_plataforma_km": 3.5},
    {"trecho": "SP-330 km 17 a 22", "rodovia": "SP-330", "km_inicial": 17, "km_final": 22, "area_verde_m2": 97000,
     "pct_manual": 35, "pct_mecanizada": 65, "equipamento_predominante": "Trator 4.7m",
     "eps_barreira_m": 1200, "defensa_m": 2100, "placas_qtd": 30, "calcada_m2": 350,
     "aceiro_km": 5, "drenagem_plataforma_km": 5, "drenagem_fora_plataforma_km": 4},
]

ROXO = "5E22F3"
ROXO_ESCURO = "3D1A9E"
LAVANDA = "EDE8FB"
LAVANDA_CLARO = "F7F4FE"
_BORDA = Border(*(Side(style="thin", color="D9D0F5") for _ in range(4)))


def _estilo_cabecalho(celula, obrigatoria: bool):
    celula.font = Font(bold=True, color="FFFFFF" if obrigatoria else ROXO_ESCURO, size=11)
    celula.fill = PatternFill("solid", fgColor=ROXO if obrigatoria else "D9CCFB")
    celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celula.border = _BORDA


def gerar_planilha_modelo() -> bytes:
    """Gera a planilha .xlsx modelo para o usuário preencher quando não há
    mapeamento Quasar. A 1ª aba ('Inventario') é a que o app lê: cabeçalho
    colorido (roxo = obrigatória, lilás = opcional), comentários em cada
    coluna, validação de dados (números ≥ 0, percentuais de 0 a 100, lista de
    equipamentos), linhas pré-formatadas e painel congelado. A aba
    'Instruções' explica o passo a passo e a aba 'Exemplo' mostra 2 trechos
    preenchidos (o app NÃO lê essas duas abas)."""
    colunas = COLUNAS_MINIMAS + COLUNAS_OPCIONAIS
    wb = Workbook()

    # ---- Aba 1: Inventario (a que o app lê) -------------------------------
    ws = wb.active
    ws.title = "Inventario"
    ws.row_dimensions[1].height = 34
    for j, c in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=j, value=c)
        _estilo_cabecalho(cel, c in COLUNAS_MINIMAS)
        cel.comment = Comment(
            f"{'OBRIGATÓRIA' if c in COLUNAS_MINIMAS else 'Opcional'} — {DESCRICAO_COLUNAS.get(c, '')}", "DimConservação"
        )
        ws.column_dimensions[get_column_letter(j)].width = 34 if c == "trecho" else (26 if c == "equipamento_predominante" else 18)
    ultima_linha = 300
    for i in range(2, ultima_linha + 1):
        for j, c in enumerate(colunas, start=1):
            cel = ws.cell(row=i, column=j)
            cel.border = _BORDA
            if i % 2 == 1:
                cel.fill = PatternFill("solid", fgColor=LAVANDA_CLARO)
            if c not in COLUNAS_TEXTO:
                cel.number_format = "#,##0.0#" if c.endswith(("_km", "km_inicial", "km_final")) or c.startswith("pct_") else "#,##0"
    ws.freeze_panes = "B2"

    def _validacao(nome_coluna, tipo, **kw):
        j = colunas.index(nome_coluna) + 1
        letra = get_column_letter(j)
        dv = DataValidation(type=tipo, allow_blank=True, showErrorMessage=True, **kw)
        ws.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{ultima_linha}")

    for c in colunas:
        if c in COLUNAS_TEXTO:
            continue
        if c.startswith("pct_"):
            _validacao(c, "decimal", operator="between", formula1="0", formula2="100",
                       errorTitle="Percentual inválido", error="Informe um número de 0 a 100.")
        else:
            _validacao(c, "decimal", operator="greaterThanOrEqual", formula1="0",
                       errorTitle="Valor inválido", error="Informe um número maior ou igual a zero.")
    nomes_eq = ",".join(n for n in pp.ROCADA_MECANIZADA)
    _validacao("equipamento_predominante", "list", formula1=f'"{nomes_eq}"',
               errorTitle="Equipamento fora da lista", error="Escolha um equipamento da lista (ou deixe em branco).",
               errorStyle="warning")

    # ---- Aba 2: Instruções --------------------------------------------------
    wi = wb.create_sheet("Instruções")
    wi.sheet_view.showGridLines = False
    wi["A1"] = "Como preencher a planilha de inventário"
    wi["A1"].font = Font(bold=True, size=16, color=ROXO_ESCURO)
    passos = [
        "1. Preencha a aba 'Inventario' (a primeira): uma linha por trecho, começando na linha 2. Não altere o cabeçalho.",
        "2. Colunas em ROXO são obrigatórias; em LILÁS são opcionais (deixe em branco o que não tiver).",
        "3. pct_manual + pct_mecanizada devem somar 100 em cada trecho. Sem elas, o app considera 100% da área em cada modalidade.",
        "4. Salve e envie o arquivo de volta na aba Inventário do app. A aba 'Exemplo' mostra 2 trechos preenchidos.",
        "5. As abas 'Instruções' e 'Exemplo' não são lidas pelo app — pode deixá-las no arquivo.",
    ]
    for i, txt in enumerate(passos, start=3):
        wi.cell(row=i, column=1, value=txt).alignment = Alignment(wrap_text=True, vertical="top")
        wi.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        wi.row_dimensions[i].height = 30
    linha0 = 9
    for j, h in enumerate(["Coluna", "Obrigatória?", "Tipo", "O que informar"], start=1):
        cel = wi.cell(row=linha0, column=j, value=h)
        _estilo_cabecalho(cel, True)
    for i, c in enumerate(colunas, start=linha0 + 1):
        obrig = c in COLUNAS_MINIMAS
        vals = [c, "Sim" if obrig else "Não", "Texto" if c in COLUNAS_TEXTO else ("% (0–100)" if c.startswith("pct_") else "Número"),
                DESCRICAO_COLUNAS.get(c, "")]
        for j, v in enumerate(vals, start=1):
            cel = wi.cell(row=i, column=j, value=v)
            cel.border = _BORDA
            cel.alignment = Alignment(wrap_text=True, vertical="center")
            if j == 2:
                cel.font = Font(bold=True, color="FFFFFF" if obrig else ROXO_ESCURO)
                cel.fill = PatternFill("solid", fgColor=ROXO if obrig else "D9CCFB")
                cel.alignment = Alignment(horizontal="center", vertical="center")
            elif i % 2 == 0:
                cel.fill = PatternFill("solid", fgColor=LAVANDA_CLARO)
        wi.row_dimensions[i].height = 32
    for letra, larg in zip("ABCD", [30, 15, 12, 95]):
        wi.column_dimensions[letra].width = larg

    # ---- Aba 3: Exemplo -----------------------------------------------------
    we = wb.create_sheet("Exemplo")
    for j, c in enumerate(colunas, start=1):
        _estilo_cabecalho(we.cell(row=1, column=j, value=c), c in COLUNAS_MINIMAS)
        we.column_dimensions[get_column_letter(j)].width = 34 if c == "trecho" else (26 if c == "equipamento_predominante" else 18)
    for i, linha in enumerate(_EXEMPLO_MODELO, start=2):
        for j, c in enumerate(colunas, start=1):
            cel = we.cell(row=i, column=j, value=linha.get(c))
            cel.border = _BORDA
            if i % 2 == 1:
                cel.fill = PatternFill("solid", fgColor=LAVANDA_CLARO)
    we.freeze_panes = "B2"
    we["A5"] = "Exemplo fictício, só para ilustrar o preenchimento — este conteúdo NÃO é lido pelo app."
    we["A5"].font = Font(italic=True, color="777777")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def validar_colunas(df: pd.DataFrame) -> list[str]:
    """Retorna a lista de colunas minimas que estao faltando (vazia = ok)."""
    faltando = [c for c in COLUNAS_MINIMAS if c not in df.columns]
    return faltando


def carregar_inventario(arquivo) -> pd.DataFrame:
    """Le um .xlsx ou .csv de inventario e devolve o DataFrame normalizado."""
    nome = getattr(arquivo, "name", str(arquivo)).lower()
    if nome.endswith(".csv"):
        df = pd.read_csv(arquivo)
    else:
        df = pd.read_excel(arquivo)

    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
    return df.dropna(how="all")


def preparar_inventario(df: pd.DataFrame) -> pd.DataFrame:
    """Calcula colunas derivadas (extensao do trecho) e garante tipos numericos.

    Se o DataFrame ja vier com `extensao_km` calculada (ex: agregacao Quasar,
    que nao soma o comprimento bruto dos poligonos para nao inflar a extensao
    com canteiro lateral/central e os dois sentidos), essa coluna e respeitada
    em vez de recalculada pela diferenca simples km_final - km_inicial.

    Se o inventario indica `equipamento_predominante` mas nao traz a area por
    equipamento (colunas mec_m2__*), a area mecanizada de cada trecho e
    atribuida por inteiro ao equipamento indicado.
    """
    df = df.copy()
    df["trecho"] = df["trecho"].astype(str)
    if "extensao_km" not in df.columns:
        df["extensao_km"] = (
            pd.to_numeric(df["km_final"], errors="coerce") - pd.to_numeric(df["km_inicial"], errors="coerce")
        ).abs()

    colunas_mec = [c for c in df.columns if str(c).startswith("mec_m2__")]
    numericas = ["km_inicial", "km_final", "area_verde_m2", "extensao_km", "largura_media_m", "inclinacao_media"] + [
        c for c in COLUNAS_OPCIONAIS if c not in COLUNAS_TEXTO
    ] + colunas_mec
    for c in numericas:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0 if c not in ("largura_media_m", "inclinacao_media") else float("nan"))
    for c in COLUNAS_TEXTO:
        if c in df.columns and c != "trecho":
            df[c] = df[c].fillna("").astype(str).str.strip()

    if "equipamento_predominante" in df.columns and not colunas_mec:
        df["equipamento_predominante"] = df["equipamento_predominante"].apply(
            lambda t: pp.canonizar_equipamento(t) if str(t).strip() else ""
        )
        area_mec = area_por_modalidade(df, "mecanizada")
        for nome in sorted({n for n in df["equipamento_predominante"] if n}):
            df[pp.coluna_area_equipamento(nome)] = area_mec.where(df["equipamento_predominante"] == nome, 0.0)
    return df


def area_por_modalidade(df: pd.DataFrame, modalidade: str) -> pd.Series:
    """Area verde de cada trecho atendida por uma modalidade ('manual' ou
    'mecanizada'). Se o inventario tiver a coluna pct_<modalidade> (vinda do
    mapeamento Quasar ou preenchida manualmente), pondera a area por ela;
    senao, assume 100% (cenario isolado, sem divisao real conhecida)."""
    coluna_pct = f"pct_{modalidade}"
    if coluna_pct in df.columns:
        return df["area_verde_m2"] * (df[coluna_pct] / 100)
    return df["area_verde_m2"]


def gerar_inventario_exemplo() -> pd.DataFrame:
    """Planilha de inventario ficticia, so para demonstracao/piloto."""
    dados = [
        {"trecho": "Trecho 1 - km 0 a km 15", "rodovia": "SP-000", "km_inicial": 0, "km_final": 15,
         "area_verde_m2": 180000, "eps_barreira_m": 3200, "defensa_m": 5400,
         "placas_qtd": 90, "bueiros_qtd": 22, "calcada_m2": 4200,
         "pct_manual": 55, "pct_mecanizada": 45, "equipamento_predominante": "Trator 1.7m"},
        {"trecho": "Trecho 2 - km 15 a km 32", "rodovia": "SP-000", "km_inicial": 15, "km_final": 32,
         "area_verde_m2": 260000, "eps_barreira_m": 5100, "defensa_m": 8300,
         "placas_qtd": 140, "bueiros_qtd": 35, "calcada_m2": 6100,
         "pct_manual": 40, "pct_mecanizada": 60, "equipamento_predominante": "Trator 4.7m"},
        {"trecho": "Trecho 3 - km 32 a km 50", "rodovia": "SP-000", "km_inicial": 32, "km_final": 50,
         "area_verde_m2": 310000, "eps_barreira_m": 4700, "defensa_m": 9600,
         "placas_qtd": 160, "bueiros_qtd": 41, "calcada_m2": 5800,
         "pct_manual": 70, "pct_mecanizada": 30, "equipamento_predominante": "Trator com braço articulado"},
    ]
    return preparar_inventario(pd.DataFrame(dados))
